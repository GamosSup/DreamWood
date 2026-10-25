#!/usr/bin/env python3
"""Ротация из Excel → картинка для приложения.
Использование: python3 tools/rot_image.py файл.xlsx rotations/2026-10-26-K.png
Нужны: libreoffice (soffice), poppler (pdftoppm), ImageMagick (convert), openpyxl."""
import sys,os,subprocess,tempfile,shutil,openpyxl
from openpyxl.worksheet.properties import PageSetupProperties
src,dst=sys.argv[1],sys.argv[2]
tmp=tempfile.mkdtemp()
wb=openpyxl.load_workbook(src);ws=wb.worksheets[-1] if len(wb.worksheets)==1 else wb.active
last=max((c.row for r in ws.iter_rows() for c in r if c.value is not None),default=ws.max_row)
lastc=max((c.column for r in ws.iter_rows() for c in r if c.value is not None),default=ws.max_column)
from openpyxl.utils import get_column_letter as L
ws.print_area=f'A1:{L(lastc)}{last}'  # без пустых строк сетки внизу
ws.page_setup.orientation='landscape';ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=1
ws.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True)
m=ws.page_margins;m.left=m.right=m.top=m.bottom=0.2
f=os.path.join(tmp,'r.xlsx');wb.save(f)
subprocess.run(['soffice','--headless','--convert-to','pdf','--outdir',tmp,f],check=True,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL,timeout=180)
subprocess.run(['pdftoppm','-r','230','-png','-f','1','-l','1',os.path.join(tmp,'r.pdf'),os.path.join(tmp,'p')],check=True)
p=[x for x in os.listdir(tmp) if x.startswith('p') and x.endswith('.png')][0]
# обрезаем белые поля и пустые строки сетки внизу
subprocess.run(['convert',os.path.join(tmp,p),'-trim','+repage','-bordercolor','white','-border','12',dst],check=True)
shutil.rmtree(tmp)
print('ok',dst)
